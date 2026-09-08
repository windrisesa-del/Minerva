from __future__ import annotations

import hashlib
import mimetypes
import re
from pathlib import Path
from typing import Any

import fitz
from docx import Document
from PIL import Image

from .assignment_import import UPLOADS_DIR


SUPPORTED_SUFFIXES = {".txt", ".docx", ".pdf", ".png", ".jpg", ".jpeg"}
FORMULA_PATTERN = re.compile(r"(\$[^$\r\n]+\$|\\\([^\r\n]+?\\\))")


def parse_uploaded_document(storage_path: str) -> dict[str, Any]:
    source = _resolve_upload(storage_path)
    suffix = source.suffix.lower()
    if suffix not in SUPPORTED_SUFFIXES:
        raise ValueError(f"document_parse 不支持 {suffix or '无扩展名'} 文件")
    if suffix == ".txt":
        return _parse_txt(source, storage_path)
    if suffix == ".docx":
        return _parse_docx(source, storage_path)
    if suffix == ".pdf":
        return _parse_pdf(source, storage_path)
    return _parse_image(source, storage_path)


def _resolve_upload(storage_path: str) -> Path:
    if not storage_path.startswith("/uploads/"):
        raise ValueError("document_parse 只允许读取 /uploads/ 下的文件")
    # Build from path parts instead of trusting platform separators from input.
    candidate = UPLOADS_DIR.joinpath(*storage_path.removeprefix("/uploads/").split("/")).resolve()
    root = UPLOADS_DIR.resolve()
    if candidate != root and root not in candidate.parents:
        raise ValueError("document_parse 文件路径越界")
    if not candidate.is_file():
        raise ValueError("document_parse 找不到上传文件")
    return candidate


def _base(source: Path, storage_path: str, format_name: str) -> dict[str, Any]:
    return {
        "schema_version": "minerva-normalized-document/0.1",
        "source": {
            "file_name": source.name,
            "path": storage_path,
            "format": format_name,
            "sha256": hashlib.sha256(source.read_bytes()).hexdigest(),
        },
        "pages": [],
        "assets": {},
        "warnings": [],
    }


def _parse_txt(source: Path, storage_path: str) -> dict[str, Any]:
    result = _base(source, storage_path, "txt")
    raw = source.read_bytes()
    text = None
    for encoding in ("utf-8-sig", "utf-16", "gb18030"):
        try:
            text = raw.decode(encoding)
            break
        except UnicodeDecodeError:
            continue
    if text is None:
        raise ValueError("TXT 编码无法识别")
    result["pages"] = [{
        "page": 1,
        "width": None,
        "height": None,
        "blocks": _split_text_blocks(text, "p1", None),
    }]
    return result


def _split_text_blocks(text: str, page_prefix: str, bbox: list[float] | None) -> list[dict[str, Any]]:
    blocks: list[dict[str, Any]] = []
    for fragment in FORMULA_PATTERN.split(text):
        if not fragment:
            continue
        block_id = f"{page_prefix}_b{len(blocks) + 1}"
        if FORMULA_PATTERN.fullmatch(fragment):
            latex = fragment[1:-1] if fragment.startswith("$") else fragment[2:-2]
            blocks.append({"id": block_id, "type": "formula", "latex": latex.strip(), "bbox": bbox})
        elif fragment.strip():
            blocks.append({"id": block_id, "type": "text", "text": fragment, "bbox": bbox})
    return blocks or [{"id": f"{page_prefix}_b1", "type": "text", "text": "", "bbox": bbox}]


def _asset_path(source: Path, label: str, suffix: str) -> tuple[Path, str]:
    digest = hashlib.sha256(source.read_bytes()).hexdigest()[:16]
    folder = UPLOADS_DIR / "adapter-assets" / digest
    folder.mkdir(parents=True, exist_ok=True)
    target = folder / f"{label}{suffix}"
    return target, f"/uploads/adapter-assets/{digest}/{target.name}"


def _parse_image(source: Path, storage_path: str) -> dict[str, Any]:
    result = _base(source, storage_path, source.suffix.lower().lstrip("."))
    with Image.open(source) as image:
        width, height = image.size
    asset_id = "asset_001"
    result["assets"][asset_id] = {
        "path": storage_path,
        "mime_type": mimetypes.guess_type(source.name)[0] or "application/octet-stream",
        "page": 1,
        "bbox": [0, 0, width, height],
        "source_path": storage_path,
    }
    result["pages"] = [{
        "page": 1,
        "width": width,
        "height": height,
        "blocks": [{"id": "p1_b1", "type": "image", "asset_id": asset_id, "bbox": [0, 0, width, height]}],
    }]
    result["warnings"].append("OCR backend is not configured; the image is preserved losslessly as an asset")
    return result


def _parse_docx(source: Path, storage_path: str) -> dict[str, Any]:
    result = _base(source, storage_path, "docx")
    document = Document(source)
    blocks: list[dict[str, Any]] = []
    sequence = 1
    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        if text:
            for block in _split_text_blocks(text, "p1", None):
                block["id"] = f"p1_b{sequence}"
                blocks.append(block)
                sequence += 1
    for table in document.tables:
        rows = [[cell.text.strip() for cell in row.cells] for row in table.rows]
        blocks.append({"id": f"p1_b{sequence}", "type": "table", "data": rows, "bbox": None})
        sequence += 1
    for index, shape in enumerate(document.inline_shapes, start=1):
        relationship_id = shape._inline.graphic.graphicData.pic.blipFill.blip.embed
        part = document.part.related_parts[relationship_id]
        suffix = Path(part.partname).suffix or ".png"
        target, asset_url = _asset_path(source, f"docx_image_{index:03d}", suffix)
        target.write_bytes(part.blob)
        asset_id = f"asset_{index:03d}"
        result["assets"][asset_id] = {
            "path": asset_url,
            "mime_type": part.content_type,
            "page": 1,
            "bbox": None,
            "source_path": storage_path,
        }
        blocks.append({"id": f"p1_b{sequence}", "type": "image", "asset_id": asset_id, "bbox": None})
        sequence += 1
    result["pages"] = [{"page": 1, "width": None, "height": None, "blocks": blocks}]
    if not blocks:
        result["warnings"].append("DOCX contained no extractable paragraphs, tables, or inline images")
    return result


def _parse_pdf(source: Path, storage_path: str) -> dict[str, Any]:
    result = _base(source, storage_path, "pdf")
    document = fitz.open(source)
    asset_index = 1
    try:
        for page_index, page in enumerate(document, start=1):
            blocks: list[dict[str, Any]] = []
            raw = page.get_text("dict")
            for block in raw.get("blocks", []):
                bbox = [round(float(value), 3) for value in block.get("bbox", [])]
                if block.get("type") == 0:
                    lines = []
                    for line in block.get("lines", []):
                        text = "".join(span.get("text", "") for span in line.get("spans", []))
                        if text.strip():
                            lines.append(text)
                    text = "\n".join(lines).strip()
                    if text:
                        for text_block in _split_text_blocks(text, f"p{page_index}", bbox):
                            text_block["id"] = f"p{page_index}_b{len(blocks)+1}"
                            blocks.append(text_block)
                elif block.get("type") == 1 and len(bbox) == 4:
                    asset_id = f"asset_{asset_index:03d}"
                    target, asset_url = _asset_path(source, f"p{page_index}_visual_{asset_index:03d}", ".png")
                    pixmap = page.get_pixmap(matrix=fitz.Matrix(2, 2), clip=fitz.Rect(bbox), alpha=False)
                    pixmap.save(target)
                    result["assets"][asset_id] = {
                        "path": asset_url,
                        "mime_type": "image/png",
                        "page": page_index,
                        "bbox": bbox,
                        "source_path": storage_path,
                    }
                    blocks.append({"id": f"p{page_index}_b{len(blocks)+1}", "type": "image", "asset_id": asset_id, "bbox": bbox})
                    asset_index += 1
            if not any(block["type"] == "text" for block in blocks):
                asset_id = f"asset_{asset_index:03d}"
                target, asset_url = _asset_path(source, f"p{page_index}_scan", ".png")
                page.get_pixmap(matrix=fitz.Matrix(2, 2), alpha=False).save(target)
                bbox = [0, 0, round(page.rect.width, 3), round(page.rect.height, 3)]
                result["assets"][asset_id] = {
                    "path": asset_url,
                    "mime_type": "image/png",
                    "page": page_index,
                    "bbox": bbox,
                    "source_path": storage_path,
                }
                blocks.append({"id": f"p{page_index}_b{len(blocks)+1}", "type": "image", "asset_id": asset_id, "bbox": bbox})
                result["warnings"].append(f"Page {page_index} has no native text; preserved as a scanned-page asset")
                asset_index += 1
            result["pages"].append({
                "page": page_index,
                "width": round(page.rect.width, 3),
                "height": round(page.rect.height, 3),
                "blocks": blocks,
            })
    finally:
        document.close()
    return result
