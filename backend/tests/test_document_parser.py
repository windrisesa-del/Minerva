from __future__ import annotations

import io
import shutil
import uuid

import fitz
from docx import Document
from fastapi.testclient import TestClient
from PIL import Image, ImageDraw
from sqlalchemy import text

from app.assignment_import import UPLOADS_DIR
from app.db import SessionLocal
from app.document_parser import parse_uploaded_document
from app.main import app


def _write(name: str, content: bytes) -> str:
    folder = UPLOADS_DIR / "adapter-test" / uuid.uuid4().hex
    folder.mkdir(parents=True, exist_ok=True)
    (folder / name).write_bytes(content)
    return f"/uploads/{folder.relative_to(UPLOADS_DIR).as_posix()}/{name}"


def _png() -> bytes:
    output = io.BytesIO()
    image = Image.new("RGB", (240, 120), "white")
    draw = ImageDraw.Draw(image)
    draw.rectangle((20, 25, 220, 95), outline="black", width=3)
    draw.line((20, 60, 220, 60), fill="black", width=2)
    image.save(output, format="PNG")
    return output.getvalue()


def test_document_parse_txt_docx_native_pdf_scan_and_image() -> None:
    txt = parse_uploaded_document(_write("paper.txt", b"Q1: solve $x^2 = 4$\nAnswer: x=2"))
    assert txt["pages"][0]["blocks"][0]["type"] == "text"
    assert any(block["type"] == "formula" for block in txt["pages"][0]["blocks"])
    assert txt["source"]["sha256"]

    docx_buffer = io.BytesIO()
    document = Document()
    document.add_paragraph("Q1: calculate 2 + 2")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Answer"
    table.cell(0, 1).text = "4"
    document.save(docx_buffer)
    docx = parse_uploaded_document(_write("paper.docx", docx_buffer.getvalue()))
    assert {block["type"] for block in docx["pages"][0]["blocks"]} == {"text", "table"}

    native_buffer = io.BytesIO()
    native = fitz.open()
    page = native.new_page()
    page.insert_text((72, 72), "Q1: Solve x^2 = 4")
    native_buffer.write(native.tobytes())
    native.close()
    parsed_pdf = parse_uploaded_document(_write("paper.pdf", native_buffer.getvalue()))
    assert any(block["type"] == "text" for block in parsed_pdf["pages"][0]["blocks"])

    scan = fitz.open()
    page = scan.new_page(width=240, height=120)
    page.insert_image(page.rect, stream=_png())
    scan_path = _write("scan.pdf", scan.tobytes())
    scan.close()
    parsed_scan = parse_uploaded_document(scan_path)
    assert parsed_scan["assets"]
    assert any("scanned-page asset" in warning for warning in parsed_scan["warnings"])

    parsed_image = parse_uploaded_document(_write("circuit.png", _png()))
    image_block = parsed_image["pages"][0]["blocks"][0]
    assert image_block["type"] == "image"
    assert image_block["asset_id"] in parsed_image["assets"]


def test_normalized_assessment_persists_as_ungraded_and_rebuilds_questions() -> None:
    with TestClient(app) as client:
        with SessionLocal() as session:
            class_id = session.execute(text(
                """
                SELECT enrollment.class_id
                FROM enrollments enrollment
                JOIN students student ON student.id = enrollment.student_id
                WHERE student.student_number = '1' AND enrollment.left_at IS NULL
                LIMIT 1
                """
            )).scalar_one_or_none()
        if class_id is None:
            return
        imported = client.post(
            "/api/assignments/import",
            data={"title": "Adapter E2E 测试_请忽略", "class_id": str(class_id)},
            files=[
                ("files", ("题目.txt", "Q1: 2+2=?".encode(), "text/plain")),
                ("files", ("1.txt", "4".encode(), "text/plain")),
            ],
        )
        if imported.status_code == 400 and "没有匹配到任何学生" in imported.text:
            return
        assert imported.status_code == 200, imported.text
        body = imported.json()
        assignment_id = body["assignment_id"]
        try:
            assert client.get(f"/api/assignments/{assignment_id}").json()["assignment"]["status"] == "draft"
            source = next(item for item in body["adapter_sources"] if item["role"] == "assessment_material")
            student_source = next(item for item in body["adapter_sources"] if item["role"] == "student_submission")
            normalized = client.post("/api/document/parse", json={"path": source["storage_key"]})
            assert normalized.status_code == 200, normalized.text
            block = normalized.json()["pages"][0]["blocks"][0]
            assessment = {
                "schema_version": "minerva-assessment/0.1",
                "status": "ungraded",
                "metadata": {"title": "Adapter E2E 测试_请忽略", "subject": "数学", "total_score": 10},
                "questions": [{
                    "question_id": "Q1",
                    "position": 1,
                    "question_type": "objective",
                    "content": [{"type": "text", "text": "2+2=?"}],
                    "reference_solution": {"answer": "4", "reasoning": ["整数加法"], "max_score": 10, "scoring_criteria": [{"score": 10, "requirement": "答案为 4"}], "partial_credit": []},
                    "analysis": {"subject": "数学", "knowledge_domain": "代数", "question_type": "基础计算", "main_concepts": ["整数加法"], "expected_path": ["直接计算"], "dependencies": [], "difficulty": "easy", "required_abilities": ["计算"]},
                    "source_references": [{"path": source["storage_key"], "page": 1, "block_ids": [block["id"]]}],
                }],
                "assets": {},
                "student_submissions": [{
                    "student_id": student_source["student_id"],
                    "normalized_documents": [student_source["storage_key"]],
                    "assets": {},
                    "answers": [{
                        "question_id": "Q1",
                        "status": "answered",
                        "content": [{"type": "text", "text": "4"}],
                        "selected_options": [],
                        "source_references": [{"path": student_source["storage_key"], "page": 1, "block_ids": ["p1_b1"]}],
                    }],
                    "uncertainties": [],
                }],
                "normalized_documents": [source["storage_key"]],
                "uncertainties": [],
            }
            saved = client.post(f"/api/assignments/{assignment_id}/assessment", json=assessment)
            assert saved.status_code == 200, saved.text
            assert saved.json()["status"] == "ungraded"
            detail = client.get(f"/api/assignments/{assignment_id}").json()
            assert detail["assignment"]["item_count"] == 1
            assert detail["assignment"]["max_score"] == 10
            items = client.get(
                "/api/minerva/read",
                params={"resource": "assignment_items", "assignment_id": assignment_id},
            ).json()["records"]
            assert items[0]["question_snapshot"]["knowledge_points"] == ["整数加法"]
            question_id = client.get(
                "/api/minerva/read",
                params={"resource": "questions", "assignment_id": assignment_id},
            ).json()["records"][0]["id"]
            answers = client.get(
                "/api/minerva/read",
                params={
                    "resource": "answer_attempts",
                    "assignment_id": assignment_id,
                    "student_id": student_source["student_id"],
                    "question_id": question_id,
                },
            ).json()["records"]
            assert answers[0]["answer_payload"]["text"] == "4"
            assert answers[0]["answer_payload"]["attachments"] == []
        finally:
            with SessionLocal.begin() as session:
                session.execute(text("DELETE FROM assignments WHERE id = CAST(:id AS uuid)"), {"id": assignment_id})


def teardown_module() -> None:
    shutil.rmtree(UPLOADS_DIR / "adapter-test", ignore_errors=True)
